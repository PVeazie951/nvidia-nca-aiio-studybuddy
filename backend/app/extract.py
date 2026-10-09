"""Text extraction from uploaded study material."""

from __future__ import annotations

import io
import subprocess
import tempfile
from pathlib import Path

IMAGE_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/gif", "image/bmp"}


def _ocr_image(raw: bytes) -> tuple[str, str]:
    """OCR an image via tesseract if available. Returns (text, error).

    Screenshots are upscaled 2x first: tesseract's accuracy on UI-sized text is
    markedly better when the glyphs are larger.
    """
    try:
        try:
            from PIL import Image

            img = Image.open(io.BytesIO(raw))
            img = img.convert("L")
            img = img.resize((img.width * 2, img.height * 2), Image.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            payload = buf.getvalue()
        except Exception:  # noqa: BLE001
            payload = raw

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as fh:
            fh.write(payload)
            tmp = fh.name
        proc = subprocess.run(
            ["tesseract", tmp, "stdout", "--psm", "6"],
            capture_output=True,
            text=True,
            timeout=180,
        )
        Path(tmp).unlink(missing_ok=True)
        if proc.returncode != 0:
            return "", f"tesseract failed: {proc.stderr.strip()[:300]}"
        return proc.stdout.strip(), ""
    except FileNotFoundError:
        return "", "tesseract not installed (OCR unavailable for screenshots)"
    except subprocess.TimeoutExpired:
        return "", "tesseract timed out"
    except Exception as exc:  # noqa: BLE001
        return "", f"ocr error: {exc}"


def _pdf_text(raw: bytes) -> tuple[str, str]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        chunks = []
        for page in reader.pages:
            try:
                chunks.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001
                continue
        text = "\n".join(chunks).strip()
        if not text:
            return "", "no selectable text found (scanned PDF? OCR not applied)"
        return text, ""
    except Exception as exc:  # noqa: BLE001
        return "", f"pdf error: {exc}"


def _docx_text(raw: bytes) -> tuple[str, str]:
    try:
        import docx

        doc = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(p for p in parts if p).strip(), ""
    except Exception as exc:  # noqa: BLE001
        return "", f"docx error: {exc}"


def _xlsx_text(raw: bytes) -> tuple[str, str]:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"# sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip(), ""
    except Exception as exc:  # noqa: BLE001
        return "", f"xlsx error: {exc}"


def extract_text(filename: str, content_type: str, raw: bytes) -> tuple[str, str]:
    """Return (extracted_text, error). Never raises."""
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    if ctype in IMAGE_TYPES or name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp")):
        return _ocr_image(raw)

    if ctype == "application/pdf" or name.endswith(".pdf"):
        return _pdf_text(raw)

    if name.endswith(".docx"):
        return _docx_text(raw)

    if name.endswith((".xlsx", ".xlsm")):
        return _xlsx_text(raw)

    if ctype.startswith("text/") or name.endswith(
        (".txt", ".md", ".json", ".yaml", ".yml", ".csv", ".log", ".py", ".sh", ".conf")
    ):
        try:
            return raw.decode("utf-8", errors="replace").strip(), ""
        except Exception as exc:  # noqa: BLE001
            return "", f"text decode error: {exc}"

    return "", f"no extractor for {content_type or name or 'unknown type'}"
